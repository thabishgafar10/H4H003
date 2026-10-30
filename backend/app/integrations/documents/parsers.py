"""
Document Parsers:
Supports PDF, DOCX, Markdown, and TXT files.
Extracts structured content with page markers and section titles.
"""

import io
import logging
import os
from typing import Any, Dict, List, Tuple

logger = logging.getLogger(__name__)


def parse_pdf(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse PDF bytes using pypdf, extracting text with page markers."""
    import pypdf

    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    pages: List[Dict[str, Any]] = []
    full_text_parts = []

    for idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        clean_text = page_text.strip()
        if clean_text:
            pages.append({
                "page_number": idx,
                "text": clean_text,
                "section_title": f"Page {idx}",
            })
            full_text_parts.append(f"--- Page {idx} ---\n{clean_text}")

    full_text = "\n\n".join(full_text_parts)
    return full_text, pages


def parse_docx(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse DOCX bytes using python-docx, preserving headings and paragraphs."""
    import docx

    doc = docx.Document(io.BytesIO(file_bytes))
    sections: List[Dict[str, Any]] = []
    full_text_parts = []
    current_heading = "General"
    current_para_group = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        if p.style.name.startswith("Heading"):
            if current_para_group:
                section_text = "\n".join(current_para_group)
                sections.append({
                    "page_number": 1,
                    "text": section_text,
                    "section_title": current_heading,
                })
                full_text_parts.append(f"## {current_heading}\n{section_text}")
                current_para_group = []
            current_heading = text
        else:
            current_para_group.append(text)

    # Process remaining paragraphs
    if current_para_group:
        section_text = "\n".join(current_para_group)
        sections.append({
            "page_number": 1,
            "text": section_text,
            "section_title": current_heading,
        })
        full_text_parts.append(f"## {current_heading}\n{section_text}")

    # Process tables if any
    for t_idx, table in enumerate(doc.tables, start=1):
        rows_text = []
        for row in table.rows:
            cell_texts = [c.text.strip().replace("\n", " ") for c in row.cells]
            rows_text.append(" | ".join(cell_texts))
        if rows_text:
            table_content = "\n".join(rows_text)
            sections.append({
                "page_number": 1,
                "text": table_content,
                "section_title": f"Table {t_idx}",
            })
            full_text_parts.append(f"### Table {t_idx}\n{table_content}")

    full_text = "\n\n".join(full_text_parts)
    return full_text, sections


def parse_text(file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """Parse plain text or Markdown bytes."""
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")

    clean_text = text.strip()
    sections = [{
        "page_number": 1,
        "text": clean_text,
        "section_title": "Full Document",
    }]
    return clean_text, sections


def parse_document(filename: str, file_bytes: bytes) -> Tuple[str, List[Dict[str, Any]]]:
    """Route document to appropriate parser based on file extension."""
    ext = os.path.splitext(filename.lower())[1]

    if ext == ".pdf":
        return parse_pdf(file_bytes)
    elif ext == ".docx":
        return parse_docx(file_bytes)
    elif ext in (".txt", ".md", ".markdown"):
        return parse_text(file_bytes)
    else:
        raise ValueError(f"No parser available for extension '{ext}'")
